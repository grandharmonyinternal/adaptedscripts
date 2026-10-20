# Builds the client copy of Billionaire Boyfriend v6 (natural spoken dialogue, Palakkad for the village family) from ep*.txt.
# Usage: python3 build_docx.py <first> <last> <out.docx>
# Client copy = title + "Episode N" heading + subheading + story paragraphs. Cue lines are dropped.
import os, re, sys, datetime
from docx import Document
from docx.shared import Pt
d = os.path.dirname(os.path.abspath(__file__))
first, last, out = int(sys.argv[1]), int(sys.argv[2]), sys.argv[3]
doc = Document()
doc.styles['Normal'].font.size = Pt(12)
doc.add_heading("BILLIONAIRE BOYFRIEND (v6 sample)", 0)
for n in range(first, last + 1):
    lines = [l.strip() for l in open(os.path.join(d, f'ep{n}.txt'), encoding='utf-8') if l.strip()]
    assert lines[0] == f'Episode {n}', lines[0]
    doc.add_heading(f'Episode {n}', 1)
    doc.add_heading(lines[1], 2)
    for l in lines[2:]:
        if re.match(r'^\[(SFX|BGM):', l):
            continue
        doc.add_paragraph(l)
now = datetime.datetime.now()
cp = doc.core_properties
cp.author = cp.last_modified_by = 'Grand Harmony Studios'
cp.title = "Billionaire Boyfriend v6 sample"
cp.created = cp.modified = now
cp.revision = 1
doc.save(out)
print('saved', out)
