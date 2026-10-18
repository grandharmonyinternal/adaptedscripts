# Builds the client copy of Paavam Paavam Kodeeswaran (Malayalam, Eps 1-43) from ep*.txt.
# Client copy = title + "Episode N" headings + story paragraphs. Cue lines ([SFX:/[BGM:) are dropped.
import os, re, sys, datetime
from docx import Document
from docx.shared import Pt
d = os.path.dirname(os.path.abspath(__file__))
out = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/Downloads/Paavam Paavam Kodeeswaran - Episodes 1-43.docx")
doc = Document()
doc.styles['Normal'].font.size = Pt(12)
doc.add_heading("PAAVAM PAAVAM KODEESWARAN", 0)
for n in range(1, 44):
    lines = [l.strip() for l in open(os.path.join(d, f'ep{n}.txt'), encoding='utf-8') if l.strip()]
    assert lines[0] == f'Episode {n}', lines[0]
    doc.add_heading(f'Episode {n}', 1)
    for l in lines[1:]:
        if re.match(r'^\[(SFX|BGM):', l):
            continue
        doc.add_paragraph(l)
now = datetime.datetime.now()
cp = doc.core_properties
cp.author = cp.last_modified_by = 'Grand Harmony Studios'
cp.title = "Paavam Paavam Kodeeswaran"
cp.created = cp.modified = now
cp.revision = 1
doc.save(out)
print('saved', out)
